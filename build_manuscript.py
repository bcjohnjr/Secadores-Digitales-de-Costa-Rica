#!/usr/bin/env python3
"""Build the four manuscript .docx files (ES/EN x named/anonymous) on the Revista Ingeniería layout.
Usage: python build_manuscript.py <template.docx> <out_dir>"""
import copy
import csv
import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import content_en
import content_es
import references
import values

HERE = Path(__file__).resolve().parent
FIGNAMES = {
    "es": {"Fig11": "Fig11_calor_por_nodo.png", "Fig13": "Fig13_esquema_nodo.png", "Fig12": "Fig12_estacion_seca.png", "Fig01": "Fig01_mapa_red.png", "Fig02": "Fig02_perfil_mensual.png", "Fig03": "Fig03_cascada_calor.png", "Fig04": "Fig04_productos_secos.png",
           "Fig05": "Fig05_economia_agricultores.png", "Fig06": "Fig06_ciclo_camarones_pargo.png", "Fig07": "Fig07_carbono.png", "Fig08": "Fig08_comparacion.png",
           "Fig09": "Fig09_entrenamiento.png", "Fig10": "Fig10_sensibilidad.png", "Fig14": "Fig14_tecnologia_secado.png", "Fig15": "Fig15_economia_comunidad.png", "FigS3": "FigS3_flujo_masa.png", "FigS8": "FigS8_validacion_era5.png"},
    "en": {"Fig11": "Fig11_heat_by_node.png", "Fig13": "Fig13_node_schematic.png", "Fig12": "Fig12_dry_season.png", "Fig01": "Fig01_network_map.png", "Fig02": "Fig02_monthly_profile.png", "Fig03": "Fig03_heat_cascade.png", "Fig04": "Fig04_dried_products.png",
           "Fig05": "Fig05_farmer_economics.png", "Fig06": "Fig06_camaron_snapper_loop.png", "Fig07": "Fig07_carbon.png", "Fig08": "Fig08_comparison.png",
           "Fig09": "Fig09_training.png", "Fig10": "Fig10_sensitivity.png", "Fig14": "Fig14_dryer_technology.png", "Fig15": "Fig15_economics_community.png", "FigS3": "FigS3_mass_flow.png", "FigS8": "FigS8_era5_validation.png"},
}
FIGWIDTH = {"Fig01": 3.5, "Fig08": 3.3, "Fig06": 5.4, "Fig10": 4.3, "Fig05": 5.8, "Fig12": 5.5, "Fig03": 5.4, "Fig04": 5.6, "Fig15": 5.8, "Fig07": 5.4, "Fig02": 5.6, "Fig13": 5.6, "Fig09": 6.2}


def S(d, name):
    for st in d.styles:
        if st.name == name:
            return st
    raise KeyError(name)


def clear_body(d):
    body = d.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)


def fmt_run(r, size=11, bold=None, italic=None):
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    r.font.size = Pt(size)
    if bold is not None:
        r.bold = bold
    if italic is not None:
        r.italic = italic


def para(d, text, style="Body Text", align=None, size=11, bold=None, spacing=2.0, italic_terms=()):
    p = d.add_paragraph(style=S(d, style))
    p.paragraph_format.line_spacing = spacing
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    if align is not None:
        p.alignment = align
    # italicise species names
    parts = [text]
    for term in italic_terms:
        new = []
        for part in parts:
            if isinstance(part, tuple):
                new.append(part); continue
            segs = part.split(term)
            for i, s in enumerate(segs):
                if s:
                    new.append(s)
                if i < len(segs) - 1:
                    new.append((term,))
        parts = new
    for part in parts:
        if isinstance(part, tuple):
            fmt_run(p.add_run(part[0]), size, bold, True)
        else:
            fmt_run(p.add_run(part), size, bold)
    return p


SPECIES = ("Lutjanus guttatus", "Litopenaeus vannamei", "Stomoxys calcitrans", "Coffea arabica", "Cosmos sulphureus", "Paulownia elongata × P. fortunei")


def add_table(d, lang, num, title, header, rows, V, widths=None, keep=False):
    cap = d.add_paragraph(style=S(d, "Body Text"))
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.line_spacing = 2.0; cap.paragraph_format.keep_with_next = True
    cap.paragraph_format.space_before = Pt(6)
    r = cap.add_run(("CUADRO " if lang == "es" else "TABLE ") + num); fmt_run(r, 11, True); r.add_break()
    fmt_run(cap.add_run(title.format(**V).upper()), 11)
    t = d.add_table(rows=len(rows) + 1, cols=len(header))
    t.style = S(d, "Table")
    small = 9 if len(header) > 4 else 10
    for i, row in enumerate([header] + rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            p.style = S(d, "Compact")
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            fmt_run(p.add_run(cell.format(**V)), small, bold=(i == 0))
            if keep:
                p.paragraph_format.keep_with_next = True  # keep this table on one page
        if True:  # rows never split across pages
            trp = t.rows[i]._tr.get_or_add_trPr()
            trp.append(trp.makeelement(qn("w:cantSplit"), {qn("w:val"): "true"}))
    if widths:
        t.autofit = False
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w * 1440)))
        tblPr = t._tbl.tblPr
        lay = tblPr.find(qn("w:tblLayout"))
        if lay is None:
            lay = tblPr.makeelement(qn("w:tblLayout"), {}); tblPr.append(lay)
        lay.set(qn("w:type"), "fixed")
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = Inches(w)
    # header row repeat
    trPr = t.rows[0]._tr.get_or_add_trPr()
    h = trPr.makeelement(qn("w:tblHeader"), {qn("w:val"): "true"}); trPr.append(h)
    sp = d.add_paragraph(style=S(d, "Body Text")); sp.paragraph_format.line_spacing = 1.0


def add_node_table(d, lang, V, num="S5"):
    rows = list(csv.DictReader((HERE / "outputs" / "nodes.csv").open(encoding="utf-8")))
    f = lambda x, k=1: values.fmt(float(x), k, lang)
    es = lang == "es"
    prodname = ({"mango": "mango", "pina": "piña", "banano": "banano", "platano": "plátano", "harinas": "harinas", "cafe": "café",
                 "harina_banano": "harina de banano", "rastrojo_pina": "rastrojo", "pulpa_cafe": "pulpa", "heno_pasto": "heno"} if es else
                {"mango": "mango", "pina": "pineapple", "banano": "banana", "platano": "plantain", "harinas": "meals", "cafe": "coffee",
                 "harina_banano": "banana meal", "rastrojo_pina": "stubble", "pulpa_cafe": "pulp", "heno_pasto": "hay"})
    body = []
    for r in rows:
        prods = [prodname[k] for k in r["products"].split(";") if k]
        if float(r["camaron_t"]) > 0:
            prods.insert(0, "camarones")
        body.append([f"{r['id']} {r['name']}", r["operator"], f(r["hydro_mw"]), f(r["it_mw"], 0), r["zone"],
                     ", ".join(prods) if prods else ("conectividad" if es else "connectivity"),
                     f(float(r["drying_heat_gwh"]) + float(r["ras_heat_gwh"])), f(r["rejected_heat_gwh"])])
    hdr = (["Nodo", "Operador", "Hidro (MW)", "TI (MW)", "Zona", "Usos del calor", "Calor usado (GWhₜ)", "Calor rechazado (GWhₜ)"] if es else
           ["Node", "Operator", "Hydro (MW)", "IT (MW)", "Zone", "Heat uses", "Heat used (GWhₜ)", "Heat rejected (GWhₜ)"])
    title = ("NODOS, CAPACIDADES Y USO DEL CALOR EN EL CASO INTERRUMPIBLE DE AÑO NORMAL. ZONAS: PN PACÍFICO NORTE, PC PACÍFICO CENTRAL, VC VALLE CENTRAL, ZN ZONA NORTE, CA CARIBE"
             if es else "NODES, CAPACITIES AND HEAT USE IN THE NORMAL-YEAR INTERRUPTIBLE CASE. ZONES: PN NORTH PACIFIC, PC CENTRAL PACIFIC, VC CENTRAL VALLEY, ZN NORTHERN ZONE, CA CARIBBEAN")
    esc = lambda t: t.replace("{", "{{").replace("}", "}}")
    add_table(d, lang, num, esc(title), hdr, [[esc(c) for c in r] for r in body], V, widths=[1.3, 1.0, 0.6, 0.45, 0.45, 1.6, 0.75, 0.75])


def add_figure(d, lang, key, caption, V, figdir):
    p = d.add_paragraph(style=S(d, "Normal")); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(str(figdir / FIGNAMES[lang][key]), width=Inches(FIGWIDTH.get(key, 6.3)))
    c = para(d, caption.format(**V), style="Body Text", align=WD_ALIGN_PARAGRAPH.CENTER, size=10, spacing=2.0)
    c.paragraph_format.space_after = Pt(8)


ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]


def resolve_numbers(body):
    """Number main tables by position and supplementary tables/figures (prefix S) by first citation in the main text."""
    def texts(b):
        if b[0] in ("p", "sp"): return [b[1]]
        if b[0] in ("fig", "sfig"): return [b[2]]
        if b[0] in ("table", "stable"): return [b[1], b[2]] + [c for row in b[4] for c in row]
        return []
    main_t = [re.match(r"§T:(\w+)§", b[1]).group(1) for b in body if b[0] == "table"]
    tmap = {k: ROMAN[i] for i, k in enumerate(main_t)}
    supp_t = [re.match(r"§T:(\w+)§", b[1]).group(1) for b in body if b[0] == "stable"] + (["nodes"] if any(b[0] == "snodetable" for b in body) else [])
    supp_f = [re.search(r"§F:(\w+)§", b[2]).group(1) for b in body if b[0] == "sfig"]
    order_t, order_f = [], []
    for b in body:
        if b[0] in ("p", "fig", "table"):
            for tx in texts(b):
                for k in re.findall(r"§T:(\w+)§", tx):
                    if k in supp_t and k not in order_t: order_t.append(k)
                for k in re.findall(r"§F:(\w+)§", tx):
                    if k in supp_f and k not in order_f: order_f.append(k)
    order_t += [k for k in supp_t if k not in order_t]; order_f += [k for k in supp_f if k not in order_f]
    tmap.update({k: f"S{i + 1}" for i, k in enumerate(order_t)})
    fmap_s = {k: f"S{i + 1}" for i, k in enumerate(order_f)}
    sub = lambda tx: re.sub(r"§F:(\w+)§", lambda m: fmap_s[m.group(1)], re.sub(r"§T:(\w+)§", lambda m: tmap[m.group(1)], tx))
    out = []
    for b in body:
        if b[0] in ("p", "sp"): out.append((b[0], sub(b[1])))
        elif b[0] in ("fig", "sfig"): out.append((b[0], b[1], sub(b[2])))
        elif b[0] in ("table", "stable"): out.append((b[0], sub(b[1]), sub(b[2]), b[3], [[sub(c) for c in row] for row in b[4]]))
        elif b[0] == "snodetable": out.append(("snodetable", tmap["nodes"]))
        else: out.append(b)
    return out


def build(template, out, lang, anonymous):
    V, R = values.values(lang)
    C = content_es if lang == "es" else content_en
    figdir = HERE / "outputs" / f"figuras_{lang}"
    d = docx.Document(template)
    clear_body(d)
    para(d, C.TITLE, style="First Paragraph", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13)
    para(d, C.TITLE_ALT, style="Body Text", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    if not anonymous:
        p = para(d, "Jonathan Karl Bily", align=WD_ALIGN_PARAGRAPH.CENTER)
        sup = p.add_run("1"); fmt_run(sup, 11); sup.font.superscript = True
        aff = ("1 Investigador independiente, Zaragoza, España. ORCID 0009-0004-3526-7271. jonathan@cloudwashlaundry.com" if lang == "es"
               else "1 Independent researcher, Zaragoza, Spain. ORCID 0009-0004-3526-7271. jonathan@cloudwashlaundry.com")
        para(d, aff, align=WD_ALIGN_PARAGRAPH.CENTER)
        para(d, ("Autor de correspondencia: Jonathan Karl Bily, jonathan@cloudwashlaundry.com" if lang == "es"
                 else "Corresponding author: Jonathan Karl Bily, jonathan@cloudwashlaundry.com"), align=WD_ALIGN_PARAGRAPH.CENTER)
    else:
        para(d, "Manuscrito anonimizado para revisión por pares" if lang == "es" else "Manuscript anonymised for peer review", align=WD_ALIGN_PARAGRAPH.CENTER)
    blocks = [("Resumen", C.ABSTRACT, C.KEYWORDS), ("Abstract", C.ABSTRACT_EN, C.KEYWORDS_EN)]
    if lang == "en":
        blocks = blocks[::-1]
    V_es, _ = values.values("es"); V_en, _ = values.values("en")
    for head, ab, kw in blocks:
        d.add_paragraph(head, style=S(d, "Heading 1"))
        para(d, ab.format(**(V_es if head == "Resumen" else V_en)), style="First Paragraph", italic_terms=SPECIES)
        para(d, kw)
    # references renumbered in order of first citation (IEEE)
    def texts(b):
        if b[0] in ("p", "sp"): return [b[1]]
        if b[0] in ("fig", "sfig"): return [b[2]]
        if b[0] in ("table", "stable"): return [b[2]] + [c for row in b[4] for c in row]
        return []
    BODY0 = resolve_numbers(C.BODY)
    seen = []
    for b in BODY0:
        for tx in texts(b):
            for m_ in re.finditer(r"\[(\d+)\]", tx):
                n = int(m_.group(1))
                if n not in seen: seen.append(n)
    allrefs = references.refs(lang, anonymous)
    missing = [i + 1 for i in range(len(allrefs)) if i + 1 not in seen]
    assert not missing, f"uncited references {missing}"
    rmap = {old_: i + 1 for i, old_ in enumerate(seen)}
    rr = lambda txt: re.sub(r"\[(\d+)\]", lambda m_: f"[{rmap[int(m_.group(1))]}]", txt)
    BODY = []
    for b in BODY0:
        if b[0] in ("p", "sp"): BODY.append((b[0], rr(b[1])))
        elif b[0] in ("fig", "sfig"): BODY.append((b[0], b[1], rr(b[2])))
        elif b[0] in ("table", "stable"): BODY.append((b[0], b[1], rr(b[2]), b[3], [[rr(c) for c in row] for row in b[4]]))
        else: BODY.append(b)
    reflist = [None] * len(allrefs)
    for old_, new_ in rmap.items():
        reflist[new_ - 1] = re.sub(r"^\[\d+\]", f"[{new_}]", allrefs[old_ - 1])
    order = [b[1] for b in BODY if b[0] == "fig"]
    fmap = {int(k[3:]): i + 1 for i, k in enumerate(order)}
    renum = lambda txt: re.sub(r"Fig\. (\d+)", lambda m: f"Fig. {fmap[int(m.group(1))]}", txt)
    for b in BODY:
        kind = b[0]
        if kind == "h1":
            d.add_paragraph(b[1], style=S(d, "Heading 1"))
        elif kind == "h2":
            d.add_paragraph(b[1], style=S(d, "Heading 2"))
        elif kind == "p":
            para(d, renum(b[1].format(**V)).replace("US$-", "−US$"), italic_terms=SPECIES)
        elif kind == "fig":
            add_figure(d, lang, b[1], renum(b[2]), V, figdir)
        elif kind == "nodetable":
            add_node_table(d, lang, V)
        elif kind == "table":
            add_table(d, lang, b[1], b[2], b[3], b[4], V, keep=len(b[4]) <= 12, widths=([1.2, 0.55, 0.6, 0.42, 0.8, 0.9, 0.68, 0.62, 0.8] if len(b[3]) == 9 and ("COMPAR" in b[2].upper()) else None))
    d.add_paragraph("Declaraciones" if lang == "es" else "Declarations", style=S(d, "Heading 1"))
    for t in (C.DECLARATIONS_ANON if anonymous else C.DECLARATIONS_NAMED):
        para(d, t)
    d.add_paragraph("Referencias" if lang == "es" else "References", style=S(d, "Heading 1"))
    for r in reflist:
        p = para(d, r, align=WD_ALIGN_PARAGRAPH.LEFT, size=11, spacing=2.0, italic_terms=SPECIES)
        p.paragraph_format.left_indent = Inches(0.35); p.paragraph_format.first_line_indent = Inches(-0.35)
    # supplementary material (separate file)
    sd = docx.Document(template)
    clear_body(sd)
    para(sd, ("Material suplementario" if lang == "es" else "Supplementary material"), style="First Paragraph", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13)
    para(sd, C.TITLE, style="Body Text", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    sps = [b for b in BODY if b[0] == "sp"]
    if sps:
        sd.add_paragraph("Notas complementarias" if lang == "es" else "Supplementary notes", style=S(sd, "Heading 1"))
        for b in sps:
            para(sd, renum(b[1].format(**V)).replace("US$-", "−US$"), italic_terms=SPECIES)
    items = []
    for b in BODY:
        if b[0] == "stable":
            items.append((0, int(b[1][1:]), b))
        elif b[0] == "snodetable":
            items.append((0, int(b[1][1:]), b))
        elif b[0] == "sfig":
            items.append((1, int(re.search(r"Fig\. S(\d+)", b[2]).group(1)), b))
    for _, _, b in sorted(items, key=lambda x: (x[0], x[1])):
        if b[0] == "stable":
            add_table(sd, lang, b[1], b[2], b[3], b[4], V)
        elif b[0] == "sfig":
            add_figure(sd, lang, b[1], renum(b[2]), V, figdir)
        elif b[0] == "snodetable":
            add_node_table(sd, lang, V, b[1])
    sd.core_properties.author = "" if anonymous else "Jonathan Karl Bily"
    sd.core_properties.title = C.TITLE
    sd.save(str(out).replace(".docx", "_Suplementario.docx" if lang == "es" else "_Supplementary.docx"))
    d.core_properties.author = "" if anonymous else "Jonathan Karl Bily"
    d.core_properties.last_modified_by = "" if anonymous else "Jonathan Karl Bily"
    d.core_properties.title = C.TITLE
    d.core_properties.comments = ""
    d.save(out)


if __name__ == "__main__":
    tpl, outdir = Path(sys.argv[1]), Path(sys.argv[2])
    outdir.mkdir(parents=True, exist_ok=True)
    for lang, anon, name in (("es", False, "Manuscrito_ES_Secadores_Digitales.docx"), ("es", True, "Manuscrito_Anonimo_ES_Secadores_Digitales.docx"),
                             ("en", False, "Manuscript_EN_Secadores_Digitales.docx"), ("en", True, "Manuscript_Anonymous_EN_Secadores_Digitales.docx")):
        build(tpl, outdir / name, lang, anon)
        print("built", name)
